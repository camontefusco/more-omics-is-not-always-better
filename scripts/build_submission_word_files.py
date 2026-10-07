from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "submission"

def setup(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(1); sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1); sec.right_margin = Inches(1)
    style = doc.styles["Normal"]
    style.font.name = "Arial"; style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.08

def add_title(doc, text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.bold = True; r.font.name = "Arial"; r.font.size = Pt(14)
    p.paragraph_format.space_after = Pt(18)

def build_cover():
    doc = Document(); setup(doc); add_title(doc, "Cover Letter — Journal of Pharmaceutical and BioTech Industry (JPBI)")
    paras = [
        "Dear Editors,",
        "We submit “When More Omics Is Not Always Better: Drug- and Dataset-Conditional Value of Multimodal Cancer Pharmacogenomics” as a Research article. The study evaluates multimodal information fusion for cancer drug-response prediction using a declared 15-compound universe spanning GDSC1, GDSC2, and CTD².",
        "The manuscript fits Journal of Pharmaceutical and BioTech Industry (JPBI) because it studies feature-level fusion in an imperfect and incomplete biomedical environment, with explicit attention to missingness, redundancy, algorithmic comparison, and cross-source transfer. The central finding is that fusion is conditional: it produced a small mean improvement across five full-universe drug-held-out splits, favored fusion in four splits and expression in one, and was slightly worse in all three leave-one-dataset-out folds.",
        "State of the art: the study is positioned against CCLE/GDSC pharmacogenomic resources, multimodal drug-response methods including MOLI, and machine-learning benchmarks of multi-omics integration [Barretina et al., 2012; Iorio et al., 2016; Tsherniak et al., 2017; Azuaje, 2017; Sharifi-Noghabi et al., 2019; Cai et al., 2022]. The contribution is evaluation design and evidence-bounded comparison, not a claim of a new deep-learning architecture.",
        "Public datasets: we used DepMap response and molecular releases for GDSC1, GDSC2, and CTD², with release identifiers and checksums recorded in the repository. Raw files are not redistributed.",
        "Declarations: the study received no external funding; the author declares no competing interests; ethics approval is not applicable because the study uses publicly available de-identified cell-line and molecular data with no human participants, patient intervention, or animal experimentation. DepMap data are acknowledged and cited according to the provider’s guidance.",
        "Validation: we report RMSE and MAE, drug-grouped held-out splits across five seeds, fold-local feature selection, leave-one-dataset-out validation, structured missingness strata, and modality redundancy summaries.",
        "Main claim and significance: adding modalities should be judged by incremental value and transferability rather than dimensionality alone. The finding that fusion fails under cross-study transfer is relevant to information-fusion systems operating across heterogeneous sources.",
        "All formal splits, feature-selection rules, derived result tables, figures, tests, and an identical reproducibility rerun are versioned in the repository. The manuscript limits claims to the declared 15-compound universe and identifies the lack of broad same-drug cross-study replication as a limitation.",
        "This manuscript has not been published previously and is not under consideration elsewhere. We are submitting a non-anonymized manuscript and do not request reviewer blinding.",
        "Sincerely,\n\nCarlos Victor Montefusco-Pereira",
    ]
    for t in paras: doc.add_paragraph(t)
    doc.save(OUT / "jpbi_cover_letter.docx")

def build_highlights():
    doc = Document(); setup(doc); add_title(doc, "Highlights")
    for t in [
        "Drug-heldout fusion gains were split-sensitive across five seeds.",
        "Fusion worsened performance in every leave-one-dataset-out fold.",
        "Structured missingness varied across expression, copy number, and mutation.",
        "Fold-local selection limits leakage in multimodal response benchmarks.",
        "More omics did not guarantee better cross-study prediction.",
    ]:
        p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(8); p.add_run(t)
    doc.save(OUT / "jpbi_highlights.docx")

if __name__ == "__main__":
    build_cover(); build_highlights()
