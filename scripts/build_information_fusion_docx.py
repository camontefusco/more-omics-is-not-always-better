from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "submission" / "information_fusion_manuscript_package.docx"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)

def set_cell(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text)); r.bold = bold; r.font.size = Pt(8)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers)); table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell(table.rows[0].cells[i], h, True); shade(table.rows[0].cells[i], "D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row): set_cell(cells[i], value)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def add_body(doc, text):
    for block in text.strip().split("\n\n"):
        if not block.strip(): continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.08
        p.add_run(block.strip())

def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p

def main():
    doc = Document()
    sec = doc.sections[0]; sec.top_margin = Inches(0.8); sec.bottom_margin = Inches(0.8); sec.left_margin = Inches(0.85); sec.right_margin = Inches(0.85)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"; styles["Normal"].font.size = Pt(10)
    for name, size, color in [("Title", 20, "17365D"), ("Heading 1", 14, "17365D"), ("Heading 2", 11, "2F5597")]:
        styles[name].font.name = "Arial"; styles[name].font.size = Pt(size); styles[name].font.color.rgb = RGBColor.from_string(color)

    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("When More Omics Is Not Always Better: Drug- and Dataset-Conditional Value of Multimodal Cancer Pharmacogenomics"); run.bold = True; run.font.size = Pt(20); run.font.color.rgb = RGBColor(23,54,93)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Carlos Victor Montefusco-Pereira\n").bold = True; p.add_run("Independent Researcher in Data Science and Artificial Intelligence in Industrial Pharmaceutics\nBerlin, Germany\nCorresponding author: [ADD EMAIL]")
    doc.add_paragraph()
    p = doc.add_paragraph(); p.add_run("Article type: ").bold = True; p.add_run("Research article")
    p = doc.add_paragraph(); p.add_run("Declarations: ").bold = True; p.add_run("Funding, competing interests, CRediT, and generative-AI declaration to be completed before submission.")
    doc.add_page_break()

    add_heading(doc, "Abstract", 1)
    add_body(doc, "Multimodal molecular profiles are often assumed to improve cancer drug-response prediction simply by adding information. We developed an evidence-bounded DepMap workflow to test that assumption across expression, copy-number, and mutation features. The prespecified evaluation universe contained 15 compounds from GDSC1, GDSC2, and CTD². Models used fold-local feature selection and Ridge regression, with drug-held-out and leave-one-dataset-out validation. In drug-held-out validation, fusion improved RMSE on two of three repeated splits, with mean RMSE 0.2005 versus 0.2063 for expression alone. In cross-study validation, fusion degraded RMSE on all three held-out datasets: CTD2 0.5649 versus 0.3276, GDSC1 0.3101 versus 0.2585, and GDSC2 0.2870 versus 0.2291. Structured missingness and weak-to-moderate row-level modality correlations indicate that additional modalities can introduce incomplete observations and domain-sensitive signal. These findings support evaluating multimodal models by incremental value, transferability, and failure modes rather than assuming that more omics is always better.")
    p = doc.add_paragraph(); p.add_run("Keywords: ").bold = True; p.add_run("information fusion; drug response; multi-omics; cancer; missing data; cross-study validation")

    add_heading(doc, "1. Introduction", 1)
    add_body(doc, "Cancer pharmacogenomics increasingly combines multiple molecular assays to predict drug response. More measurements can capture complementary biology, but they also increase dimensionality, missingness, and sensitivity to study-specific measurement processes. We therefore prespecified a workflow that treats incremental predictive value, missingness, calibration, and modality conflict as joint evaluation targets.")
    add_heading(doc, "1.1. Brief literature context", 2)
    add_body(doc, "Large cell-line resources established the empirical basis for this task. The Cancer Cell Line Encyclopedia and the Genomics of Drug Sensitivity in Cancer linked molecular profiles to pharmacologic response across large panels, while later DepMap releases expanded standardized molecular and dependency data [1–3]. These resources also show why response prediction is not purely a genomic problem: lineage and multiple molecular data types contribute to drug sensitivity associations [2].")
    add_body(doc, "Prior computational work has used classical machine learning, deep learning, and multimodal integration for drug-response prediction [4–6]. MOLI is a prominent late-integration example that combines expression, copy-number, and mutation features and reported gains in external validations [5]. However, published comparisons often differ in response metric, drug universe, feature processing, split design, and external-study transfer. The present study addresses a narrower methodological question: whether adding modalities improves a fixed, leakage-controlled baseline under drug-heldout and cross-study evaluation. It is intended as a validation benchmark, not as a competing deep-learning model.")

    add_heading(doc, "2. Methods", 1)
    add_body(doc, "We harmonized DepMap response and molecular data from GDSC1, GDSC2, and CTD² using locked release identifiers. Expression, copy-number, and damaging-mutation features were evaluated separately and in concatenated fusion models. Within each held-out fold, the top 2,000 features per modality were selected using training rows only. Models used median imputation, standardization, and Ridge regression. Drug-held-out splits grouped rows by drug; leave-one-dataset-out folds held out each source dataset in turn. Outer joins were retained for structured missingness analyses.")

    add_heading(doc, "3. Results", 1)
    add_body(doc, "Drug-held-out performance was split-sensitive: fusion improved two of three seeds and worsened slightly on one. In leave-one-dataset-out validation, expression was consistently stronger than fusion, with the largest fusion degradation on CTD2. The signed fusion-minus-expression RMSE deltas and all frozen metrics are available in the machine-readable result tables. Row-level modality summaries showed weak-to-moderate correlations, while outer-joined tables showed dataset-dependent modality availability.")
    lood = json.loads((ROOT/"data/processed/leave_one_dataset_out_results.json").read_text())
    rows=[]
    for d,f in lood["folds"].items(): rows.append([d, f["metrics"]["expression"]["rmse"], f["metrics"]["copy_number"]["rmse"], f["metrics"]["mutation"]["rmse"], f["metrics"]["fusion"]["rmse"]])
    p=doc.add_paragraph(); p.add_run("Table 1. ").bold=True; p.add_run("Leave-one-dataset-out RMSE.")
    add_table(doc,["Held-out dataset","Expression","Copy number","Mutation","Fusion"], [[r[0]]+[f"{x:.4f}" for x in r[1:]] for r in rows])
    drug_rows=[]
    for seed in ["20260821","20260822","20260823"]:
        d=json.loads((ROOT/f"data/processed/drug_heldout_results_{seed}.json").read_text()); drug_rows.append([seed, f"{d['metrics']['expression']['rmse']:.4f}", f"{d['metrics']['fusion']['rmse']:.4f}", f"{d['metrics']['fusion']['rmse']-d['metrics']['expression']['rmse']:+.4f}"])
    p=doc.add_paragraph(); p.add_run("Table 2. ").bold=True; p.add_run("Drug-held-out repeated-seed RMSE.")
    add_table(doc,["Seed","Expression","Fusion","Fusion delta"],drug_rows)
    for image, caption in [(ROOT/"outputs/figures/lood_rmse.png","Figure 1. Leave-one-dataset-out performance."),(ROOT/"outputs/figures/modality_conflict.png","Figure 2. Fusion minus expression RMSE across held-out folds.")]:
        doc.add_picture(str(image), width=Inches(6.5)); p=doc.paragraphs[-1]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; cap=doc.add_paragraph(caption); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER

    add_heading(doc, "4. Discussion", 1)
    add_body(doc, "The results do not support a universal multimodal advantage. Fusion can help within a drug-held-out universe yet fail under study transfer, consistent with domain-specific response scales, assay coverage, and feature distributions. The study is limited by the small prespecified 15-drug universe, limited same-drug overlap across datasets, and the use of compact Ridge baselines rather than a broad model class. The conclusions are therefore about evaluation design and conditional predictive value, not a claim that any modality is biologically uninformative.")
    add_heading(doc, "5. Data and code availability", 1)
    add_body(doc, "Code, manifests, derived result summaries, figures, and reproducibility notes are provided in the project repository. Raw DepMap files must be obtained from the provider under the applicable terms; see the public-data access note. The repository does not redistribute raw molecular matrices.")
    add_heading(doc, "6. References", 1)
    refs=["[1] Barretina J, et al. The Cancer Cell Line Encyclopedia enables predictive modelling of anticancer drug sensitivity. Nature. 2012;483:603–607. doi:10.1038/nature11003.","[2] Iorio F, et al. A landscape of pharmacogenomic interactions in cancer. Cell. 2016;166:740–754. doi:10.1016/j.cell.2016.06.017.","[3] Tsherniak A, et al. Defining a cancer dependency map. Cell. 2017;170:564–576.e16. doi:10.1016/j.cell.2017.06.010.","[4] Azuaje F. Computational models for predicting drug responses in cancer research. Brief Bioinform. 2017;18:820–829. doi:10.1093/bib/bbw065.","[5] Sharifi-Noghabi H, et al. MOLI: multi-omics late integration with deep neural networks for drug response prediction. Bioinformatics. 2019;35:i501–i509. doi:10.1093/bioinformatics/btz318.","[6] Cai Z, et al. Machine learning for multi-omics data integration in cancer. iScience. 2022;25:103798. doi:10.1016/j.isci.2022.103798.","[7] DepMap. How should I cite DepMap data? DepMap Portal. Updated 2025."]
    for ref in refs: doc.add_paragraph(ref)

    doc.add_page_break(); add_heading(doc,"Supplementary material",1)
    add_body(doc,"Supplementary Table S1 reports the complete/missing test-row counts for the three drug-heldout seeds. Supplementary Table S2 records the frozen artifact and validation scope. Figures are supplied separately in the submission package.")
    p=doc.add_paragraph(); p.add_run("Supplementary Table S1. ").bold=True; p.add_run("Drug-held-out missingness strata.")
    add_table(doc,["Seed","Complete rows","Missing rows","Total rows"],[["20260821","1513","1046","2559"],["20260822","1458","1271","2729"],["20260823","1530","1078","2608"]])
    p=doc.add_paragraph(); p.add_run("Supplementary Table S2. ").bold=True; p.add_run("Release scope and limitations.")
    add_table(doc,["Item","Status"],[["Evaluation universe","15 prespecified compounds across GDSC1, GDSC2, CTD²"],["Drug-heldout validation","3 persisted seeds; fold-local top-2,000 features"],["Cross-study validation","LOOD folds for CTD2, GDSC1, and GDSC2"],["MoA mapping","15/15 declared compounds; 5/316 full GDSC1 universe"],["Raw data","Provider files not redistributed; checksums recorded"]])
    OUT.parent.mkdir(exist_ok=True); doc.save(OUT); print(OUT)

if __name__ == "__main__": main()
